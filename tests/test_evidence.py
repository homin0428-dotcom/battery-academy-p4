"""근거 FMEA 데이터셋·자동 판정 엔진 테스트."""
import copy
import shutil
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import core  # noqa: E402
import evidence_engine as ee  # noqa: E402
import kb_schema as S  # noqa: E402

CFG = core.load_config()
VALS, _ = core.resolve_inputs(CFG)
VALS = {**VALS, "kb_effect_selection": {}}   # 기술군 중앙값 기준 회귀 검사 (설계 고정은 test_v122)
KB = ee.load_kb()


def _kb():
    return copy.deepcopy(KB)


def _eff(**k):
    base = {c: "" for c in S.columns("effects")}
    base.update(dict(effect_id="T-EF", lever_id="DRY-ZONE", metric="system_saving_ratio", unit="%", basis="t",
                     boundary="계통전체", chemistry="NCM", cell_format="파우치", scale="양산", source_id="R4",
                     locator="p.1", excerpt="t", evidence_grade="B", verification="원문 확인"))
    base.update({kk: str(v) for kk, v in k.items()})
    return base


def _with(*rows):
    kb = _kb()
    kb["effects"] = kb["effects"] + list(rows)
    return kb


# ------------------------------------------------------------ 데이터셋 무결성
def test_seed_dataset_has_no_errors():
    _clean, issues = ee.validate_kb(KB)
    assert [i for i in issues if i["level"] == "error"] == []


def test_every_effect_has_source_and_locator():
    clean, _ = ee.validate_kb(KB)
    for r in clean["effects"]:
        assert r["source_id"] and r["locator"], r["effect_id"]


def test_example_rows_ignored():
    kb = _with(_eff(effect_id="EX-EF-9", value_typ=50))
    clean, issues = ee.validate_kb(kb)
    assert all(r["effect_id"] != "EX-EF-9" for r in clean["effects"])


@pytest.mark.parametrize("bad,msg", [
    (dict(unit="kWh"), "단위"),
    (dict(value_low=30, value_typ=10, value_high=40), "최솟값"),
    (dict(lever_id="NOPE"), "레버ID"),
    (dict(source_id="NOPE"), "근거ID"),
    (dict(boundary="어딘가"), "허용되지 않은 값"),
    (dict(value_typ=150), "허용 범위"),
    (dict(value_typ="약 10"), "숫자가 아님"),
])
def test_validation_rejects_bad_rows(bad, msg):
    row = _eff(**{"value_typ": 10, **bad})
    clean, issues = ee.validate_kb(_with(row))
    errs = [i for i in issues if i["level"] == "error" and i["id"] == "T-EF"]
    assert errs and any(msg in e["msg"] for e in errs)
    assert all(r["effect_id"] != "T-EF" for r in clean["effects"])


def test_duplicate_id_rejected():
    clean, issues = ee.validate_kb(_with(_eff(value_typ=10), _eff(value_typ=12)))
    assert any("중복" in i["msg"] for i in issues)


# ------------------------------------------------------------ 경계 환산
def test_pair_metric_derives_saving():
    clean, _ = ee.validate_kb(_with(_eff(metric="energy_before_after", value_before=100, value_after=75, unit="kWh")))
    r = next(x for x in clean["effects"] if x["effect_id"] == "T-EF")
    assert r["metric_eff"] == "system_saving_ratio" and r["typ"] == pytest.approx(0.25)


def test_plant_basis_uses_explicit_share():
    clean, _ = ee.validate_kb(_with(_eff(value_typ=10, boundary="공장전체", boundary_share=0.4)))
    r = next(x for x in clean["effects"] if x["effect_id"] == "T-EF")
    out = ee.to_system_basis(r, clean, 0.35)
    assert out["use"] == ee.USE_CALC and out["typ"] == pytest.approx(0.25)


def test_plant_basis_falls_back_to_kb_share_then_task_share():
    # 1순위: 같은 문헌의 계통 비중(경계비중 열) - Degen 표에서 역산한 드라이룸 0.247
    clean, _ = ee.validate_kb(_kb())
    r = next(x for x in clean["effects"] if x["effect_id"] == "EF-001")   # 드라이룸 미니환경 16.1% (공장전체)
    out = ee.to_system_basis(r, clean, 0.25)
    assert "같은 문헌" in out["how"] and out["typ"] == pytest.approx(0.161 / 0.247)
    assert "타 문헌 비중 환산" not in out["flags"]
    # 2순위: 경계비중이 없으면 KB 비중 근거(타 문헌) - 별도 표시
    kb = _kb()
    for e in kb["effects"]:
        if e["effect_id"] == "EF-001":
            e["boundary_share"] = ""
    clean1, _ = ee.validate_kb(kb)
    r1 = next(x for x in clean1["effects"] if x["effect_id"] == "EF-001")
    out1 = ee.to_system_basis(r1, clean1, 0.25)
    sh = ee.share_evidence(clean1, "드라이룸")
    assert "KB 계통 비중" in out1["how"] and "타 문헌 비중 환산" in out1["flags"]
    assert out1["typ"] == pytest.approx(0.161 / sh[1])
    # 3순위: KB 비중 근거도 없으면 과제 비중으로 환산하고 '가정' 표시
    kb["effects"] = [e for e in kb["effects"] if e["lever_id"] != "BASE-DRM"]
    clean2, _ = ee.validate_kb(kb)
    r2 = next(x for x in clean2["effects"] if x["effect_id"] == "EF-001")
    out2 = ee.to_system_basis(r2, clean2, 0.25)
    assert "과제 비중" in out2["how"] and out2["typ"] == pytest.approx(0.644)


def test_lever_judgement_counts_papers_not_rows():
    """같은 논문의 여러 행(Degen NIR·레이저)은 독립 근거 1건으로 센다."""
    kb = _kb()
    kb["effects"] = [e for e in kb["effects"] if not (e["lever_id"] == "DRY-RAD" and e["source_id"] != "R1")]
    clean, _ = ee.validate_kb(kb)
    eff = ee.classify_effects(clean, [35, 25, 10, 5, 25])
    rows = ee.lever_table(clean, eff, ["DRY-RAD"], {"pending": []})
    rad = next(r for r in rows if r["lever_id"] == "DRY-RAD")
    assert rad["n_calc"] == 2 and rad["n_sources"] == 1 and rad["judge"] == "근거 약함"


def test_plant_basis_over_100pct_is_excluded():
    clean, _ = ee.validate_kb(_with(_eff(value_typ=30, boundary="공장전체", boundary_share=0.2)))
    r = next(x for x in clean["effects"] if x["effect_id"] == "T-EF")
    assert ee.to_system_basis(r, clean, 0.35)["use"] == ee.USE_EXCL


@pytest.mark.parametrize("boundary", ["복합계통", "장비단위", "계통일부"])
def test_unconvertible_boundaries_are_reference_only(boundary):
    clean, _ = ee.validate_kb(_with(_eff(value_typ=50, boundary=boundary)))
    r = next(x for x in clean["effects"] if x["effect_id"] == "T-EF")
    assert ee.to_system_basis(r, clean, 0.35)["use"] == ee.USE_REF


def test_partial_boundary_with_share_is_converted():
    clean, _ = ee.validate_kb(_with(_eff(value_typ=60, boundary="계통일부", boundary_share=0.5)))
    r = next(x for x in clean["effects"] if x["effect_id"] == "T-EF")
    out = ee.to_system_basis(r, clean, 0.35)
    assert out["use"] == ee.USE_CALC and out["typ"] == pytest.approx(0.30)


# ------------------------------------------------------------ 적용 조건
def test_chemistry_mismatch_excluded_and_unverified_excluded():
    ev = ee.evaluate(VALS, kb=_with(_eff(effect_id="T-LFP", value_typ=50, chemistry="LFP"),
                                    _eff(effect_id="T-UNV", value_typ=50, verification="미확인")))
    use = {e["effect_id"]: e["use"] for e in ev["effects"]}
    assert use["T-LFP"] == ee.USE_EXCL and use["T-UNV"] == ee.USE_EXCL
    assert ev["systems"]["건조"]["status"] == "근거 없음"   # 둘 다 제외 → 건조 근거 없음 유지


def test_scale_difference_is_conditional():
    clean, _ = ee.validate_kb(_kb())
    r = next(x for x in clean["effects"] if x["effect_id"] == "EF-019")   # Weng 파일럿
    t, reasons = ee.applicability(r)
    assert t == "조건부" and any("규모" in x for x in reasons)


# ------------------------------------------------------------ 결합 규칙
def test_same_group_not_multiplied():
    # DRY-ZONE과 DRY-AIR는 같은 그룹 → 곱하지 않고 큰 값만
    kb = _with(_eff(effect_id="T-1", lever_id="DRY-ZONE", value_typ=20), _eff(effect_id="T-2", lever_id="DRY-AIR", value_typ=10))
    ev = ee.evaluate(VALS, kb=kb)
    assert ev["systems"]["건조"]["typ"] == pytest.approx(0.20)


def test_different_groups_multiply():
    kb = _with(_eff(effect_id="T-1", lever_id="DRY-ZONE", value_typ=20), _eff(effect_id="T-2", lever_id="DRY-HP", value_typ=10))
    ev = ee.evaluate(VALS, kb=kb)
    assert ev["systems"]["건조"]["typ"] == pytest.approx(1 - 0.8 * 0.9)


def test_package_covers_components():
    kb = _with(_eff(effect_id="T-1", lever_id="DRY-ZONE", value_typ=20),
               _eff(effect_id="T-P", lever_id="DRY-PKG", value_typ=35))
    ev = ee.evaluate(VALS, kb=kb, selected=["DRY-ZONE", "DRY-PKG"])
    assert "DRY-ZONE" in ev["dropped_by_package"]
    assert ev["systems"]["건조"]["typ"] == pytest.approx(0.35)


# ------------------------------------------------------------ 판정·역산
def test_seed_verdict_and_v10_check():
    ev = ee.evaluate(VALS)
    assert ev["verdict"] == "근거상 미달"
    v10 = {r["system"]: r["status"] for r in ev["v10_check"]}
    assert v10["건조"] == "근거 없음"            # v10 건조 40%는 선택 레버 근거 없음
    assert v10["드라이룸"] == "근거 범위 안"
    assert v10["화성"] == "근거 범위 안"
    c = ev["cases"]
    assert c["conservative"]["saving"] <= c["typical"]["saving"] <= c["optimistic"]["saving"]


def test_cases_use_core_formulas_exactly():
    ev = ee.evaluate(VALS)
    t = ev["cases"]["typical"]
    pl = core.plant_energy(VALS["w"], VALS["f"], t["r"], VALS["q"], t["p_add"], "pt", VALS["target_intensity_ratio"])
    assert t["E_total"] == pytest.approx(pl["E_total"])


def test_verdict_ladder():
    big = _with(_eff(effect_id="T-1", lever_id="DRY-ZONE", value_low=60, value_high=70))
    assert ee.evaluate(VALS, kb=big)["verdict"] == "근거상 충족"
    mid = _with(_eff(effect_id="T-1", lever_id="DRY-ZONE", value_low=0, value_typ=50, value_high=60))
    assert ee.evaluate(VALS, kb=mid)["verdict"] == "조건부 충족 (대표값 기준)"


def test_pending_burden_creates_hold_and_capacity_upper_bound():
    ev = ee.evaluate(VALS)
    assert any("미정량" in h for h in ev["holds"])
    assert ev["cases"]["typical"]["capacity_upper_bound"] is True   # FOR-DATA 진단 점유 미정량


def test_quantified_burden_enters_cases():
    kb = _kb()
    for r in kb["fmea"]:
        if r["fmea_id"] == "FM-005":
            r.update(burden_low="0.1", burden_typ="0.2", burden_high="0.5", unit="h")
    ev = ee.evaluate(VALS, kb=kb)
    assert ev["cases"]["conservative"]["occupancy_h"] == pytest.approx(0.5)
    assert ev["cases"]["optimistic"]["occupancy_h"] == pytest.approx(0.1)


def test_combo_search_finds_drying_alternatives_and_skips_long_term():
    ev = ee.evaluate(VALS)
    cb = ev["combos"]
    assert "DRY-DRYCOAT" in cb["long_term_excluded"]
    assert all("DRY-DRYCOAT" not in x["add"] for x in cb["results"])
    best = cb["results"][0]
    assert "DRY-RAD" in best["add"] and best["verdict"] != "근거상 미달"


def test_time_check_against_formation_target():
    ev = ee.evaluate(VALS)
    assert ev["time"]["target_ratio"] == pytest.approx(0.8)
    assert "보수적" in ev["time"]["status"]


def test_cop_param_check_flags_mode_a_assumption():
    ev = ee.evaluate(VALS)
    p = next(x for x in ev["params"] if "COP_h" in x["param"])
    assert p["status"] == "근거 범위 밖"


def test_no_scores_generated():
    ev = ee.evaluate(VALS)
    text = repr(ev["levers"]) + repr(ev["verdict"])
    for banned in ("RPN", "CEI", "확률", "AP="):
        assert banned not in text


def test_unknown_selected_lever_rejected():
    with pytest.raises(ee.KBError):
        ee.evaluate(VALS, selected=["NOPE"])


# ------------------------------------------------------------ 엑셀 가져오기
def test_template_roundtrip_import(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    import re
    sys.path.insert(0, str(ROOT / "tools"))
    import kb_import
    import make_kb_template
    x = tmp_path / "t.xlsx"
    make_kb_template.build(x)
    wb = openpyxl.load_workbook(x)
    ws = wb["정량효과"]
    keys = [re.search(r"\((\w+)\)", c.value).group(1) for c in ws[2]]
    good = _eff(effect_id="JS-1", value_low=8, value_high=15)
    bad = _eff(effect_id="JS-2", value_typ=10, unit="kWh")
    for r in (good, bad):
        ws.append([r.get(k) or None for k in keys])
    wb.save(x)
    kbdir = tmp_path / "kb"
    shutil.copytree(ee.KB_DIR, kbdir)
    rep = kb_import.import_xlsx(x, kbdir, do_merge=True)
    assert rep["accepted"]["effects"] == 1 and [e["id"] for e in rep["errors"]] == ["JS-2"]
    ev = ee.evaluate(VALS, kb_dir=kbdir)
    assert ev["systems"]["건조"]["typ"] == pytest.approx(0.115)
    # v10채택 열이 보존되어 선택 레버가 유지되는지 (머리글 숫자 포함 key)
    assert "DRY-ZONE" in ev["selected"]


def test_old_template_without_new_columns_rejected(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    sys.path.insert(0, str(ROOT / "tools"))
    import kb_import
    import make_kb_template
    x = tmp_path / "t.xlsx"
    make_kb_template.build(x)
    wb = openpyxl.load_workbook(x)
    ws = wb["정량효과"]
    col = next(c.column for c in ws[2] if "(boundary)" in c.value)
    ws.delete_cols(col)
    wb.save(x)
    with pytest.raises(ee.KBError, match="열이 없음"):
        kb_import.read_xlsx(x)


def test_all_kb_text_renders_in_report_font():
    """PDF 폰트에 없는 글리프가 빈칸으로 사라지면 '−0.8%p'가 '0.8%p'가 된다 - 변환 후 전부 그려져야 함."""
    fontTools = pytest.importorskip("fontTools.ttLib")
    import evidence_report as E
    cmap = fontTools.TTFont(str(ROOT / "assets" / "fonts" / "NanumGothic.ttf")).getBestCmap()
    missing = set()
    for sheet, rows in KB.items():
        for r in rows:
            for v in r.values():
                missing |= {c for c in E.safe_text(v) if ord(c) > 127 and ord(c) not in cmap}
    assert not missing, f"폰트에 없는 글리프: {sorted(missing)} - evidence_report.GLYPH_MAP에 추가"
    assert E.safe_text("−0.8%p, 150 µm") == "-0.8%p, 150 μm"


# ------------------------------------------------------------ 통합 (core·보고서·CLI)
def test_run_scenario_attaches_evidence_and_keeps_v10_values():
    res = core.run_scenario(CFG)
    assert res["plant"]["saving_ratio"] == pytest.approx(0.2865, abs=1e-9)   # 기존 v10 경로 불변
    assert res["evidence"]["status"] == core.NUM_OK and res["evidence"]["conclusion"]


def test_lever_override_through_custom_scenario():
    res = core.run_scenario(CFG, "custom", {"kb_selected_levers": ["FOR-FAST", "DRM-LOCAL", "DRY-RAD", "DRY-HSC"]})
    sel = ["FOR-FAST", "DRM-LOCAL", "DRY-RAD", "DRY-HSC"]
    assert res["evidence"]["selected"] == sel
    assert res["evidence"]["verdict"] == ee.evaluate(core.resolve_inputs(CFG)[0], selected=sel)["verdict"]


def test_broken_dataset_does_not_block_main_calculation(tmp_path):
    (tmp_path / "levers.csv").write_text("lever_id\n", encoding="utf-8")
    ev = core.run_evidence(core.resolve_inputs(CFG)[0], kb_dir=tmp_path)
    assert ev["status"] == core.NUM_ERR and "근거 판정 불가" in ev["error"]


def test_evidence_report_export(tmp_path):
    import evidence_report as E
    import reporting
    res = core.run_scenario(CFG)
    out = E.export(res, tmp_path, reporting.make_run_info("test"))
    assert out["pages"] >= 3
    for n in ("evidence_report.pdf", "evidence_systems.png", "evidence_cases.png", "kb_cases.csv",
              "kb_lever_judgement.csv", "kb_effect_trace.csv", "kb_combos.csv"):
        assert (tmp_path / n).exists(), n
    from pypdf import PdfReader
    text = "".join(p.extract_text() for p in PdfReader(str(tmp_path / "evidence_report.pdf")).pages)
    assert "근거상 미달" in text and "EF-001" in text


# ------------------------------------------------------------ 외부 평가 반영 (증산·에너지 종류·전환 조건·비교표·적용 순서)
def test_production_includes_aging_slots():
    pr = ee.evaluate(VALS)["production"]
    assert pr["aging_required_slot_gain"] == pytest.approx(1.2 * 0.9 / 0.95 - 1)
    assert pr["aging_allowed_days_without_slots"] == pytest.approx(14 * 0.95 / (1.2 * 0.9))
    assert pr["verdict"] == "증산 미완결 - 에이징 보관능력 미확보"   # 슬롯 확보 근거 없음


def test_aging_slot_evidence_closes_production_gap():
    kb = _with(_eff(effect_id="T-S", lever_id="AGE-RACK", metric="slot_gain_ratio", value_low=15, value_high=20,
                    energy_carrier="무관"))
    pr = ee.evaluate(VALS, kb=kb)["production"]
    assert pr["aging_status"].startswith("근거상 확보")
    assert pr["verdict"] in ("증산 조건부", "증산 성립")   # 화성 점유 부담 미정량이면 조건부


def test_heat_based_saving_dropped_when_heating_not_electric():
    """가열이 가스·스팀이면: 전력 절대 증감 근거 우선 → 없으면 설비 전기 변화율 → 둘 다 없으면 참고 전용."""
    vals = dict(VALS, kb_electric_heating={"드라이룸": False, "건조": False})
    ev = ee.evaluate(vals)
    e = {x["effect_id"]: x for x in ev["effects"]}
    assert e["EF-044"]["use"] == ee.USE_CALC and e["EF-044"]["typ"] < 0                         # Schütte NIR 전력 +1.28
    assert e["EF-006"]["use"] == ee.USE_REF and "절대 증감 근거가 있어" in e["EF-006"]["how"]     # Degen 비율 근사는 참고
    assert e["EF-029"]["use"] == ee.USE_CALC and e["EF-029"]["typ"] == pytest.approx(0.15)      # DRM-DEW: 비율 근거만 있음
    assert e["EF-003"]["use"] == ee.USE_REF                                                     # Lechner 합산(변화율 없음)
    assert ee.lever_saving(ev["effects"], "DRY-RAD")["hi"] < 0                                  # 가스면 복사건조는 전력 증가
    assert "DRY-RAD" not in ev["combos"]["results"][0]["add"]
    vals_e = dict(VALS)                                                                         # 전기 가열이면 절대값 행은 참고
    assert next(x for x in ee.evaluate(vals_e)["effects"] if x["effect_id"] == "EF-044")["use"] == ee.USE_REF


def test_primary_energy_row_is_reference_only():
    ev = ee.evaluate(VALS)
    e = next(x for x in ev["effects"] if x["effect_id"] == "EF-010")   # Gu 2025 1차에너지
    assert e["use"] == ee.USE_REF


def test_electricity_only_plant_intensity_used_for_money():
    ev = ee.evaluate(VALS)
    pe = ev["options"]["plant"]
    assert pe["ids"] == ["EF-027"] and pe["gwh_lo"] == pytest.approx(500) and pe["gwh_hi"] == pytest.approx(650)


def test_option_money_math():
    ev = ee.evaluate(dict(VALS, electricity_price_krw_per_kwh=100.0))
    v10 = ev["options"]["rows"][0]
    assert v10["pt_saved_typ"] == pytest.approx(111.2 - 85.62)
    assert v10["money_typ"]["krw_eok"][0] == pytest.approx(25.58 / 100 * 500 * 1e6 * 100 / 1e8)
    assert v10["money_typ"]["invest_eok"][3][0] == pytest.approx(3 * v10["money_typ"]["krw_eok"][0])


def test_price_scales_money_linearly():
    a = ee.evaluate(dict(VALS, electricity_price_krw_per_kwh=100.0))["options"]["rows"][1]["money_typ"]["krw_eok"][0]
    b = ee.evaluate(dict(VALS, electricity_price_krw_per_kwh=180.0))["options"]["rows"][1]["money_typ"]["krw_eok"][0]
    assert b == pytest.approx(a * 1.8)


def test_flip_thresholds_reverse_correctly():
    ev = ee.evaluate(VALS)
    t = next(x for x in ev["thresholds"] if x["kind"] == "근거 없는 레버의 기준값" and x["system"] == "건조")
    # 그 값을 건조 근거로 넣으면 정확히 대표 기준 경계(격차 0)에 도달해야 함
    kb = _with(_eff(effect_id="T-1", lever_id="DRY-ZONE", value_typ=t["value"] * 100, energy_carrier="전력"))
    gap = ee.evaluate(VALS, kb=kb)["cases"]["typical"]["gap"]
    assert gap == pytest.approx(0.0, abs=1e-9)


def test_roadmap_stages_by_change_type():
    rm = ee.evaluate(VALS)["roadmap"]
    st1 = {l["lever_id"] for l in rm[0]["levers"]}
    st2 = {l["lever_id"] for l in rm[1]["levers"]}
    assert "FOR-FAST" in st1 and "AGE-EARLY" in st1        # 운영변경 → 1년차
    assert "DRY-RAD" in st2 and "DRM-LOCAL" in st2          # 설비개조 → 2년차
    assert all(l["gates"] for l in rm[0]["levers"] if l["lever_id"] == "FOR-FAST")


def test_conclusion_leads_with_recommendation():
    ev = ee.evaluate(VALS)
    txt = ee.conclusion(ev)
    assert txt.startswith("[종합]") and "[추천]" in txt and "[증산]" in txt and "[추천 전환 조건]" in txt


def test_previous_template_without_optional_columns_still_imports(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    sys.path.insert(0, str(ROOT / "tools"))
    import kb_import
    import make_kb_template
    x = tmp_path / "t.xlsx"
    make_kb_template.build(x)
    wb = openpyxl.load_workbook(x)
    ws = wb["정량효과"]
    col = next(c.column for c in ws[2] if "(energy_carrier)" in c.value)   # 선택 열 삭제 = 템플릿 v2
    ws.delete_cols(col)
    wb.save(x)
    kb = kb_import.read_xlsx(x)
    assert "effects" in kb
