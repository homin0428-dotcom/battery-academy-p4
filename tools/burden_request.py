"""미정량 관리 부담 추정 요청서(엑셀) 만들기.

추천 조합의 FMEA 행 중 부담지표는 있는데 값이 비어 있는 행을 모아, 팀이 범위(최솟값~최댓값)를 채우도록 한다.
맨 위에 '허용 예산'을 같이 보여준다: 이 합을 넘으면 판정이 바뀐다.
채운 값은 팀 작업 파일(근거FMEA 시트)의 같은 FMEA_ID 행 부담 칸에 옮겨 적으면 된다 (추정이면 근거등급 E).

사용: python tools/burden_request.py [출력.xlsx]
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

import core  # noqa: E402
import evidence_engine as ee  # noqa: E402
import kb_schema as S  # noqa: E402

F = "맑은 고딕"


def build(out):
    v, _ = core.resolve_inputs(core.load_config())
    ev = ee.evaluate(v)
    clean, rec = ev["clean"], ev["recommended"]["levers"]
    B = [wi * (fi + (1 - fi) * v["q"]) for wi, fi in zip(v["w"], v["f"])]
    _s, _d, bur, _c = ee._run_cases(clean, ev["effects"], rec, v, B)
    fm = {r["fmea_id"]: r for r in clean["fmea"]}
    pr = ev["production_rec"]
    q, cu, pu, days = v["q"], v["current_utilization"], v["planned_utilization"], pr["aging_days"]
    ratio = (pr.get("aging_ratio") or (1, 1, 1))[1]
    gain = (pr.get("aging_gain") or (0, 0, 0))[1]
    age_allow = days * (1 + gain) * pu / (q * cu) - ratio * days
    p_allow = next((t["value"] for t in ev["thresholds"] if t["kind"] == "관리 부담 허용치"), None)
    h_allow = next((t["value"] for t in ev["thresholds"] if t["kind"] == "화성 관리 부담 허용치"), None)

    wb = Workbook()
    ws = wb.active
    ws.title = "추정요청"
    ws.append(["미정량 관리 부담 추정 요청 (추천 조합: " + " + ".join(ev["recommended"]["add"] + ev["recommended"].get("production_add", [])) + ")"])
    ws["A1"].font = Font(name=F, size=13, bold=True)
    ws.append(["값이 비어 있어 지금은 0으로 계산 중인 부담입니다. 아는 범위만 최솟값~최댓값으로 적어 주세요. 근거가 없으면 팀 추정(등급 E)으로 적고 이유를 메모에 씁니다."])
    ws.append([])
    ws.append(["허용 예산(이 합을 넘으면 판정이 바뀜)", "", "", ""])
    ws.cell(ws.max_row, 1).font = Font(name=F, bold=True)
    budget = [("추가 전력 합(added_power_ratio × 해당 계통 전력)", f"{p_allow:.2f} pt (기준 공장 100 대비)" if p_allow is not None else "-",
               "넘으면 추천 조합이 대표 기준 미달"),
              ("화성 추가 채널 점유 합(added_occupancy_h)", f"{h_allow:.2f} h" if h_allow is not None else "-", "넘으면 화성 처리능력 1.2배 미달"),
              ("에이징 추가 보관일 합(added_aging_days)", f"{age_allow:.2f} 일",
               f"현재 에이징 근거 기간비 {ratio:.2f}(대표) 기준. 넘으면 슬롯 부족")]
    for b in budget:
        ws.append(list(b))
    ws.append([])
    head = ["FMEA_ID", "레버", "계통", "고장모드", "부담지표", "단위", "최솟값(입력)", "최댓값(입력)", "근거·이유(입력)",
            "근거등급(입력: 근거 있으면 A~D, 추정이면 E)", "참고: 계통 전력 B(pt)"]
    ws.append(head)
    hr = ws.max_row
    for c in ws[hr]:
        c.font, c.fill = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3B4D")
    for p in bur["pending"]:
        r = fm[p["fmea_id"]]
        sys_ = ee._lever(clean, r["lever_id"])["system"]
        unit = S.BURDEN_METRICS[p["metric"]]["canon"]
        ws.append([p["fmea_id"], r["lever_id"], sys_, r["failure_mode"], S.BURDEN_METRICS[p["metric"]]["ko"], unit, "", "", "", "",
                   round(B[ee.SYS_IDX[sys_]], 2) if p["metric"] == "added_power_ratio" else ""])
    for col, wdt in zip("ABCDEFGHIJK", (14, 12, 8, 40, 30, 8, 12, 12, 40, 18, 14)):
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.font.name != F:
                c.font = Font(name=F, size=10, bold=c.font.bold, color=c.font.color)
    ws.append([])
    ws.append(["예: added_power_ratio 0.02 = 그 계통 전력의 2% 추가. 드라이룸 B가 30 pt면 0.6 pt. "
               "added_aging_days 0.5 = 셀당 평균 반나절 더 보관."])
    wb.save(out)
    return out, len(bur["pending"])


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "outputs" / "관리부담_추정요청.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    print(build(out))
