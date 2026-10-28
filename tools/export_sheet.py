"""병합된 데이터셋을 팀 공유용 엑셀(구글 시트 업로드용)로 내보낸다.

- 입력 시트(근거자료·정량효과·근거FMEA): 팀 입력층(data/kb_team)만 채운다 → 팀원이 이어서 고치고 다시 보내도
  Claude 조사분이 팀 층으로 복사되어 굳지 않는다.
- '전체_근거자료', '전체_정량효과', '전체_근거FMEA' 시트: 병합된 마스터 전체(읽기 전용 참고). kb_import는 이 시트를 읽지 않는다.
- '판정_요약' 시트: 현재 데이터로 돌린 종합 판정.

사용: python tools/export_sheet.py [출력.xlsx]
"""
import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
import make_kb_template as T  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

import core  # noqa: E402
import evidence_engine as ee  # noqa: E402
import kb_schema as S  # noqa: E402


def _read(path):
    return list(csv.DictReader(open(path, encoding="utf-8-sig"))) if Path(path).exists() else []


def build(out_path, style_rows=15):
    T.STYLE_ROWS = style_rows
    levers = _read(ROOT / "data" / "kb" / "levers.csv")
    team = {s: _read(ROOT / "data" / "kb_team" / f"{s}.csv") for s in ("sources", "effects", "fmea")}
    master = {s: _read(ROOT / "data" / "kb" / f"{s}.csv") for s in ("sources", "effects", "fmea")}
    wb = Workbook()
    T._guide(wb)
    refs = T._lists_sheet(wb, levers)
    for s in ("sources", "effects", "fmea"):
        T._data_sheet(wb, s, refs, T.EXAMPLES[s] + team[s])
    T._data_sheet(wb, "levers", refs, levers)
    # 전체 데이터(참고)
    team_ids = {s: {r[S.ID_COL[s]] for r in team[s]} for s in team}
    for s, title in (("sources", "전체_근거자료"), ("effects", "전체_정량효과"), ("fmea", "전체_근거FMEA")):
        ws = wb.create_sheet(title)
        cols = S.columns(s)
        ws.cell(1, 1, f"{title} - 병합된 전체 데이터(읽기 전용 참고, 가져오기 때 무시됨). 수정은 '{S.SHEET_KO[s]}' 시트에서").font = T.F_H2
        ws.append(["입력층"] + [S.header(ko, k) for k, ko, _r, _d in S.SHEETS[s]])
        for c in ws[2]:
            c.font, c.fill = T.F_HEAD, T.FILL_HEAD_OPT
        for r in master[s]:
            ws.append(["팀 입력" if r[S.ID_COL[s]] in team_ids[s] else "Claude 조사"] + [r.get(k, "") for k in cols])
        ws.freeze_panes = "C3"
    # 판정 요약
    v, _ = core.resolve_inputs(core.load_config())
    ev = ee.evaluate(v)
    ws = wb.create_sheet("판정_요약")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 110
    ws.append(["현재 데이터 기준 자동 판정", ""])
    ws["A1"].font = T.F_TITLE
    rec = ev["recommended"]
    rows = [("종합 판정(추천 조합)", ev["overall"]["label"]),
            ("추천 추가 레버", " + ".join(rec["add"] + (rec.get("production_add") or [])) or "없음")]
    rows += [(a["axis"], a["status"]) for a in ev["overall"]["axes"]]
    rows += [("v10 레버만", f"대표 {ev['cases']['typical']['saving'] * 100:.1f}% ({ev['verdict']})"),
             ("데이터 규모", f"자료 {ev['counts']['sources']} · 정량효과 {ev['counts']['effects']} · FMEA {ev['counts']['fmea']}"),
             ("결론", ee.conclusion(ev))]
    for a, b in rows:
        ws.append([a, b])
        ws.cell(ws.max_row, 1).font = Font(name=T.FONT, size=10, bold=True)
        ws.cell(ws.max_row, 2).alignment = T.WRAP
    order = ["작성가이드", "판정_요약", "근거자료", "정량효과", "근거FMEA", "레버",
             "전체_근거자료", "전체_정량효과", "전체_근거FMEA", "목록"]
    wb._sheets = [wb[n] for n in order]
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "outputs" / "근거FMEA_병합본_구글시트용.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    print(build(out))
